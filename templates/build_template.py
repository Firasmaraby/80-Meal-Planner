#!/usr/bin/env python3
"""Builds MealPlan_Template.xlsx for the Eighty meal-plan importer.

Run:  python3 templates/build_template.py
Reads src/data/foods.json and src/data/swapGroups.json when present (re-run after they change).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent / "MealPlan_Template.xlsx"
FOODS_JSON = ROOT / "src" / "data" / "foods.json"
SWAP_JSON = ROOT / "src" / "data" / "swapGroups.json"

CUISINES = ["Mediterranean", "North American", "Asian fusion", "Middle Eastern", "Mexican",
            "Indian", "Japanese", "Italian", "Caribbean", "West African", "Other"]
SLOTS = ["breakfast", "lunch", "dinner", "snack"]

HEADER_FILL = PatternFill("solid", fgColor="1F6F54")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
EXAMPLE_FONT = Font(italic=True, color="808080", name="Calibri")
EXAMPLE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D0D0D0")
HEADER_BORDER = Border(bottom=Side(style="medium", color="0F3D2E"))


def load_json(path):
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:  # tolerate a half-written file from another agent
        print(f"warning: could not parse {path.name}: {exc}")
        return None


def as_list(data, *keys):
    """Accept [..] or {"foods": [..]} / {key: {...}} shapes."""
    if data is None:
        return None
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for k in keys:
            if isinstance(data.get(k), list):
                return data[k]
        return [dict(v, key=v.get("key", k)) if isinstance(v, dict) else {"key": k, "foodIds": v}
                for k, v in data.items()]
    return None


def style_header(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = HEADER_BORDER
        c.alignment = Alignment(vertical="center", horizontal="left")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def write_example(ws, values, wrap_cols=()):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=2, column=i, value=v)
        c.font = EXAMPLE_FONT
        c.fill = EXAMPLE_FILL
        c.alignment = Alignment(vertical="top", wrap_text=(i in wrap_cols))


def add_dv(ws, dv, col, last_row):
    dv.add(f"{col}2:{col}{last_row}")
    ws.add_data_validation(dv)


def build_instructions(wb):
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 95

    rows = [
        ("title", "Eighty meal plan template", None),
        ("sub", "Fill in the sheets, save the file, and upload it in the Eighty admin page. "
                "Nutrition is calculated automatically, so you never type it.", None),
        ("gap", None, None),
        ("h", "The big idea", None),
        ("kv", "Nutrition is automatic", "Never type calories, protein or vitamins. You only say WHICH foods and HOW MANY GRAMS. "
                                         "Eighty looks the foods up in its cited food database and calculates every number, so all plans stay consistent."),
        ("kv", "Grey italic rows", "The first row in each sheet is an example (its id starts with EXAMPLE). The importer skips it. "
                                   "Copy it for a quick start, or delete it."),
        ("gap", None, None),
        ("h", "The sheets", None),
        ("kv", "Plans", "One row per meal plan: id, name, calorie range, cuisine, tags and a short description. "
                        "calorie_min and calorie_max (1000 to 5000) set the calorie band the plan is listed under; every day should fall inside it."),
        ("kv", "Meals", "One row per meal. plan_id says which plan it belongs to, day (1 to 14) and slot (breakfast, lunch, dinner, snack) say when. "
                        "Give each meal a unique meal_id. A day can have several snacks; they keep the order you list them in."),
        ("kv", "Ingredients", "One row per ingredient. meal_id says which meal it belongs to; food_id (or food_name) says what it is; grams says how much."),
        ("kv", "Foods", "Read-only reference list of every food in the app: id, name, category. Use Ctrl+F or the filter arrows to find a Food ID."),
        ("kv", "SwapGroups", "Read-only list of swap group keys (for example leafy greens) and the foods inside them."),
        ("gap", None, None),
        ("h", "How to reference a food", None),
        ("kv", "Option 1: Food ID (best)", "Look the food up on the Foods sheet and copy its id into the food_id column. This is exact and always wins."),
        ("kv", "Option 2: plain name", "Leave food_id empty and type a plain name in food_name, e.g. 'rolled oats'. The importer fuzzy-matches it against the food database. "
                                       "A confident match is accepted but flagged with a warning showing what it picked, so you can check it. "
                                       "If it is not sure, it reports an error with up to 3 suggestions. Copy the right id into food_id and re-upload."),
        ("kv", "grams", "Weight of the ingredient as used (a number above 0, decimals allowed, e.g. 37.5). Use the weight that goes into the meal, not the shop pack size."),
        ("gap", None, None),
        ("h", "Swap groups", None),
        ("kv", "What they do", "A swap group lets the person using the app swap one ingredient for another in the same family "
                               "(for example spinach for kale in 'leafy greens') while keeping the meal balanced."),
        ("kv", "How to use", "In the Ingredients sheet, pick a key from the swap_group dropdown (keys come from the SwapGroups sheet). "
                             "Leave it empty for ingredients that should not be swapped."),
        ("gap", None, None),
        ("h", "Rules the importer checks", None),
        ("kv", "Required", "Plans: plan_id, plan_name, calorie_min, calorie_max, cuisine. Meals: plan_id, day, slot, meal_id, meal_name. Ingredients: meal_id, grams, and food_id or food_name."),
        ("kv", "IDs", "plan_id and meal_id must be unique. Every plan_id in Meals must exist in Plans, and every meal_id in Ingredients must exist in Meals."),
        ("kv", "Completeness", "Every plan needs at least one meal, and every meal at least one ingredient. calorie_min must be lower than calorie_max."),
        ("kv", "Calorie check", "After import, each day's calculated calories are compared with the plan's calorie band. Days outside it get a warning (not an error)."),
        ("kv", "Cuisine", "Pick from the dropdown or type your own; anything is accepted."),
        ("kv", "tags", "Comma-separated, lowercase works best: vegetarian, pescatarian, dairy-free, gluten-free, high-protein."),
        ("kv", "version", "A whole number. Bump it when you re-upload a changed plan (leave empty for 1)."),
        ("kv", "Do not", "rename columns, add or reorder sheets, or merge cells in the data sheets. Extra blank rows are fine."),
    ]
    r = 2
    for kind, a, b in rows:
        if kind == "title":
            c = ws.cell(row=r, column=2, value=a)
            c.font = Font(bold=True, size=20, color="1F6F54")
            ws.row_dimensions[r].height = 32
        elif kind == "sub":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            c = ws.cell(row=r, column=2, value=a)
            c.font = Font(size=12, color="404040")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 36
        elif kind == "h":
            for col in (2, 3):
                cc = ws.cell(row=r, column=col)
                cc.fill = HEADER_FILL
            c = ws.cell(row=r, column=2, value=a)
            c.font = Font(bold=True, size=12, color="FFFFFF")
            ws.row_dimensions[r].height = 22
        elif kind == "kv":
            k = ws.cell(row=r, column=2, value=a)
            k.font = Font(bold=True)
            k.alignment = Alignment(vertical="top", wrap_text=True)
            v = ws.cell(row=r, column=3, value=b)
            v.alignment = Alignment(vertical="top", wrap_text=True)
            for col in (2, 3):
                ws.cell(row=r, column=col).border = Border(bottom=THIN)
            ws.row_dimensions[r].height = max(18, 15 * (len(b) // 88 + 1) + 3)
        r += 1


def build():
    wb = Workbook()
    build_instructions(wb)

    foods = as_list(load_json(FOODS_JSON), "foods")
    groups = as_list(load_json(SWAP_JSON), "swapGroups", "groups")

    # ---------------- Plans
    ws = wb.create_sheet("Plans")
    headers = ["plan_id", "plan_name", "calorie_min", "calorie_max", "cuisine", "tags", "description", "author", "version"]
    style_header(ws, headers, [18, 32, 13, 13, 18, 30, 60, 20, 10])
    write_example(ws, ["EXAMPLE-plan", "Example: Mediterranean 1800", 1700, 1900, "Mediterranean",
                       "pescatarian, high-protein", "Replace or delete this row. Rows starting with EXAMPLE are ignored.",
                       "Your name", 1], wrap_cols=(7,))
    dv = DataValidation(type="whole", operator="between", formula1="1000", formula2="5000", allow_blank=True,
                        showErrorMessage=True, errorTitle="Calories", error="Enter a whole number between 1000 and 5000.",
                        showInputMessage=True, promptTitle="Calories", prompt="Whole number, 1000 to 5000")
    add_dv(ws, dv, "C", 500)
    dv = DataValidation(type="whole", operator="between", formula1="1000", formula2="5000", allow_blank=True,
                        showErrorMessage=True, errorTitle="Calories", error="Enter a whole number between 1000 and 5000.",
                        showInputMessage=True, promptTitle="Calories", prompt="Whole number, 1000 to 5000")
    add_dv(ws, dv, "D", 500)
    dv = DataValidation(type="list", formula1='"' + ",".join(CUISINES) + '"', allow_blank=True,
                        showErrorMessage=False, errorStyle="warning", showInputMessage=True,
                        promptTitle="Cuisine", prompt="Pick from the list or type your own")
    add_dv(ws, dv, "E", 500)
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                        showErrorMessage=True, errorTitle="Version", error="Whole number, 1 or higher.")
    add_dv(ws, dv, "I", 500)

    # ---------------- Meals
    ws = wb.create_sheet("Meals")
    headers = ["plan_id", "day", "slot", "meal_id", "meal_name", "prep_minutes", "instructions", "photo_url"]
    style_header(ws, headers, [18, 8, 14, 18, 34, 14, 60, 36])
    write_example(ws, ["EXAMPLE-plan", 1, "breakfast", "EXAMPLE-meal", "Example: Oats with berries", 10,
                       "Simmer the oats, top with berries.", ""], wrap_cols=(7,))
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="14", allow_blank=True,
                        showErrorMessage=True, errorTitle="Day", error="Day must be a whole number from 1 to 14.",
                        showInputMessage=True, promptTitle="Day", prompt="1 to 14")
    add_dv(ws, dv, "B", 5000)
    dv = DataValidation(type="list", formula1='"' + ",".join(SLOTS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Slot", error="Pick breakfast, lunch, dinner or snack.")
    add_dv(ws, dv, "C", 5000)
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="Prep time", error="Whole minutes, 0 or more.")
    add_dv(ws, dv, "F", 5000)

    # ---------------- SwapGroups (built before Ingredients validation refers to it)
    wg = wb.create_sheet("SwapGroups")
    style_header(wg, ["key", "label", "food_ids"], [26, 32, 100])
    n_groups = 0
    if groups:
        for g in groups:
            ids = g.get("foodIds", g.get("food_ids", [])) or []
            wg.append([g.get("key", ""), g.get("label", ""), ", ".join(str(i) for i in ids)])
            n_groups += 1
    else:
        wg.cell(row=2, column=2, value="Note: src/data/swapGroups.json was not found when this template was built. "
                                       "Re-run templates/build_template.py to fill this list.").font = Font(italic=True, color="C00000")
    wg.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False)
    wg.sheet_properties.tabColor = "999999"

    # ---------------- Ingredients
    ws = wb.create_sheet("Ingredients", index=3)
    headers = ["meal_id", "food_id", "food_name", "grams", "swap_group", "note"]
    style_header(ws, headers, [18, 22, 34, 10, 22, 30])
    example_food_id = foods[0].get("id", "") if foods else ""
    example_food_name = foods[0].get("name", "") if foods else "Rolled oats"
    write_example(ws, ["EXAMPLE-meal", example_food_id, example_food_name, 50, "", "dry weight"])
    dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="Grams", error="Grams must be a number greater than 0.",
                        showInputMessage=True, promptTitle="Grams", prompt="Weight in grams, above 0 (decimals OK)")
    add_dv(ws, dv, "D", 20000)
    last = max(n_groups + 30, 60)
    dv = DataValidation(type="list", formula1=f"SwapGroups!$A$2:$A${last}", allow_blank=True,
                        showErrorMessage=True, errorStyle="warning", errorTitle="Swap group",
                        error="That key is not on the SwapGroups sheet. Continue anyway?",
                        showInputMessage=True, promptTitle="Swap group", prompt="Optional. Pick a key from SwapGroups")
    add_dv(ws, dv, "E", 20000)

    # ---------------- Foods
    wf = wb.create_sheet("Foods")
    style_header(wf, ["id", "name", "category"], [26, 60, 18])
    if foods:
        for f in foods:
            wf.append([f.get("id", ""), f.get("name", ""), f.get("category", "")])
        wf.auto_filter.ref = f"A1:C{len(foods) + 1}"
    else:
        wf.cell(row=2, column=2, value="Note: src/data/foods.json was not found when this template was built. "
                                       "Re-run templates/build_template.py to fill this list.").font = Font(italic=True, color="C00000")
    wf.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False)
    wf.sheet_properties.tabColor = "999999"

    # sheet order: Instructions, Plans, Meals, Ingredients, Foods, SwapGroups
    wb.move_sheet("Foods", offset=-(wb.sheetnames.index("Foods") - 4))
    for name, color in (("Plans", "1F6F54"), ("Meals", "1F6F54"), ("Ingredients", "1F6F54"), ("Instructions", "E0A100")):
        wb[name].sheet_properties.tabColor = color

    wb.save(OUT)
    print(f"wrote {OUT} (foods: {len(foods) if foods else 'none'}, swap groups: {n_groups or 'none'})")


if __name__ == "__main__":
    build()

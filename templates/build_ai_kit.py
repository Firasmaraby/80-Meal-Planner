#!/usr/bin/env python3
"""Builds the AI meal-plan kit from src/data/foods.everyday.json (the Everyday tier) and AI_Plan_Builder_Brief.md.

Run:  python3 templates/build_ai_kit.py
Run it after templates/build_template.py regenerates the template (that rebuild drops these sheets).

Writes:
  templates/AI_Food_List.csv          non-rare foods, nutrients per 100 g, sorted by category then name
  templates/MealPlan_Template.xlsx    adds "AI Brief" and "Targets" sheets right after "Instructions"
  templates/Sample_Plans_Filled.xlsx  same two sheets, if the file exists
"""
import csv
import json
import math
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

TEMPLATES = Path(__file__).resolve().parent
sys.path.insert(0, str(TEMPLATES))
from foods_bundle import load_everyday  # noqa: E402
ROOT = TEMPLATES.parent
BRIEF_MD = TEMPLATES / "AI_Plan_Builder_Brief.md"
FOOD_CSV = TEMPLATES / "AI_Food_List.csv"
TEMPLATE_XLSX = TEMPLATES / "MealPlan_Template.xlsx"
SAMPLE_XLSX = TEMPLATES / "Sample_Plans_Filled.xlsx"

CSV_HEADER_LINE = "# Per 100 g. Empty = not measured. Use food_id exactly."
CSV_HEADER = ("food_id,name,category,serving,kcal,protein_g,carbs_g,fat_g,sat_fat_g,fiber_g,"
              "vit_a_ug,vit_c_mg,vit_d_ug,vit_e_mg,vit_k_ug,folate_ug,b12_ug,choline_mg,"
              "calcium_mg,iron_mg,magnesium_mg,potassium_mg,zinc_mg,selenium_ug,omega3_g").split(",")

# CSV column -> key in the food's "n" dict (omega3_g is computed separately)
NUTRIENT_MAP = {
    "kcal": "kcal", "protein_g": "protein", "carbs_g": "carbs", "fat_g": "fat",
    "sat_fat_g": "satFat", "fiber_g": "fiber", "vit_a_ug": "vitA", "vit_c_mg": "vitC",
    "vit_d_ug": "vitD", "vit_e_mg": "vitE", "vit_k_ug": "vitK", "folate_ug": "folate",
    "b12_ug": "b12", "choline_mg": "choline", "calcium_mg": "calcium", "iron_mg": "iron",
    "magnesium_mg": "magnesium", "potassium_mg": "potassium", "zinc_mg": "zinc",
    "selenium_ug": "selenium",
}
OMEGA3_KEYS = ("ala", "epa", "dha")

HEADER_FILL = PatternFill("solid", fgColor="1F6F54")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
TARGET_HEADERS = ["Nutrient", "Pass at (per day)", "Unit", "Good sources"]


def fmt(x):
    """Round to at most 2 decimals, drop trailing zeros. None (or NaN) -> empty cell."""
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return ""
    v = round(float(x), 2)
    if v == 0:
        return "0"
    s = f"{v:.2f}".rstrip("0").rstrip(".")
    return s


def serving_text(food):
    servings = food.get("servings") or []
    if not servings:
        return ""
    first = servings[0]
    label, grams = first.get("label"), first.get("g")
    if label is None or grams is None:
        return ""
    return f"{label} ({fmt(grams)} g)"


def omega3(n):
    present = [n[k] for k in OMEGA3_KEYS if n.get(k) is not None]
    return sum(present) if present else None


def food_row(food):
    n = food.get("n") or {}
    values = {
        "food_id": food["id"],
        "name": food["name"],
        "category": food["category"],
        "serving": serving_text(food),
        "omega3_g": fmt(omega3(n)),
    }
    for col, key in NUTRIENT_MAP.items():
        values[col] = fmt(n.get(key))
    return [values[h] for h in CSV_HEADER]


def write_food_csv(foods):
    kept = [f for f in foods if not f.get("rare")]
    kept.sort(key=lambda f: (f["category"], f["name"].lower(), f["id"]))
    with FOOD_CSV.open("w", encoding="utf-8", newline="") as fh:
        fh.write(CSV_HEADER_LINE + "\n")
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(CSV_HEADER)
        for food in kept:
            writer.writerow(food_row(food))
    return len(kept)


def parse_targets(md_text):
    """Rows of the 'Nutrition targets' markdown table as (nutrient, amount, unit, sources)."""
    lines = md_text.splitlines()
    start = next((i for i, line in enumerate(lines)
                  if line.strip().lower().startswith("## nutrition targets")), None)
    if start is None:
        raise ValueError("no '## Nutrition targets' section in the brief")

    rows, in_table = [], False
    for line in lines[start + 1:]:
        s = line.strip()
        if s.startswith("|"):
            in_table = True
            cells = [c.strip() for c in s.split("|")[1:-1]]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                continue  # |---|---| separator
            if cells[0] == "Nutrient":
                continue  # header row
            if len(cells) != 3:
                raise ValueError(f"expected 3 cells in targets row: {s}")
            name, pass_text, sources = cells
            m = re.fullmatch(r"([\d,]+(?:\.\d+)?)\s+(.+)", pass_text)
            if not m:
                raise ValueError(f"cannot read amount and unit from {pass_text!r}")
            number = float(m.group(1).replace(",", ""))
            number = int(number) if number.is_integer() else number
            rows.append((name, number, m.group(2).strip(), sources))
        elif in_table or s.startswith("#"):
            break
    if not rows:
        raise ValueError("no rows found in the Nutrition targets table")
    return rows


def install_sheets(wb, lines, targets):
    for name in ("AI Brief", "Targets"):
        if name in wb.sheetnames:
            del wb[name]
    pos = wb.sheetnames.index("Instructions") + 1

    ws = wb.create_sheet("AI Brief", pos)
    ws.column_dimensions["A"].width = 110
    for i, line in enumerate(lines, start=1):
        if not line:
            continue
        cell = ws.cell(row=i, column=1)
        cell.value = line
        cell.data_type = "s"  # keep lines such as "=== PLANS ===" as text, not formulas
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    wt = wb.create_sheet("Targets", pos + 1)
    for i, title in enumerate(TARGET_HEADERS, start=1):
        cell = wt.cell(row=1, column=i, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
    for r, (name, amount, unit, sources) in enumerate(targets, start=2):
        wt.cell(row=r, column=1, value=name).alignment = Alignment(vertical="top")
        wt.cell(row=r, column=2, value=amount).alignment = Alignment(vertical="top", horizontal="left")
        wt.cell(row=r, column=3, value=unit).alignment = Alignment(vertical="top")
        wt.cell(row=r, column=4, value=sources).alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in zip("ABCD", (24, 18, 14, 80)):
        wt.column_dimensions[col].width = width
    wt.row_dimensions[1].height = 22
    wt.freeze_panes = "A2"


def save_in_place(wb, path):
    tmp = path.with_name(path.stem + ".tmp" + path.suffix)
    wb.save(tmp)
    os.replace(tmp, path)


def main():
    foods = load_everyday()
    brief_text = BRIEF_MD.read_text(encoding="utf-8")
    targets = parse_targets(brief_text)
    lines = brief_text.splitlines()

    n_csv = write_food_csv(foods)
    print(f"wrote {FOOD_CSV} ({n_csv} foods)")
    print(f"brief: {len(lines)} lines; targets: {len(targets)} rows")

    for path, required in ((TEMPLATE_XLSX, True), (SAMPLE_XLSX, False)):
        if not path.exists():
            if required:
                sys.exit(f"missing {path}")
            print(f"skipped {path.name} (not found)")
            continue
        wb = load_workbook(path)
        install_sheets(wb, lines, targets)
        wb.active = 0
        save_in_place(wb, path)
        print(f"updated {path} sheets: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()

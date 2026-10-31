"""Fill MealPlan_Template.xlsx with the bundled sample plans → Sample_Plans_Filled.xlsx (a worked example)."""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'templates'))
from foods_bundle import load_everyday  # noqa: E402
plans = json.loads((ROOT / 'src/data/plans.json').read_text())
foods = {f['id']: f['name'] for f in load_everyday()}
wb = load_workbook(ROOT / 'templates/MealPlan_Template.xlsx')


def append(ws_name, rows):
    ws = wb[ws_name]
    headers = [c.value for c in ws[1]]
    for row in rows:
        ws.append([row.get(h) for h in headers])


append('Plans', [
    {'plan_id': p['id'], 'plan_name': p['name'], 'calorie_min': p['band']['min'], 'calorie_max': p['band']['max'],
     'cuisine': p['cuisine'], 'tags': ', '.join(p['tags']), 'description': p.get('description'), 'author': p.get('author'),
     'version': p['version']}
    for p in plans])
append('Meals', [
    {'plan_id': p['id'], 'day': d['day'], 'slot': m['slot'], 'meal_id': m['id'], 'meal_name': m['name'],
     'prep_minutes': m.get('prepMinutes'), 'instructions': m.get('instructions'), 'photo_url': m.get('photoUrl')}
    for p in plans for d in p['days'] for m in d['meals']])
append('Ingredients', [
    {'meal_id': m['id'], 'food_id': i['foodId'], 'food_name': foods.get(i['foodId']), 'grams': i['grams'],
     'swap_group': i.get('swapGroup'), 'note': i.get('note')}
    for p in plans for d in p['days'] for m in d['meals'] for i in m['ingredients']])
out = ROOT / 'templates/Sample_Plans_Filled.xlsx'
wb.save(out)
print('wrote', out)

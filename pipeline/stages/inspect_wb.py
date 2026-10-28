"""--inspect: print each sheet's name, row count, first rows and columns, plus suggested nutrient mappings.

Uses openpyxl in read-only/streaming mode so a 70 MB workbook is scanned without loading it into memory.
"""
from __future__ import annotations

import json
import re
from typing import Any

from .ingest import norm_col

# (regex on norm_col(column), NutrientKey)
SUGGEST: list[tuple[str, str]] = [
    (r"^(energ|energy).*kcal|^kcal|energykcal|^energy$", "kcal"),
    (r"(energ|energy).*kj|^kj$", "kcal"),
    (r"^protein|proteintotal", "protein"),
    (r"^lipid|^totalfat|^fat(g|total)?$|^fattotal", "fat"),
    (r"^carbo|^carbs|availablecarb|carbohydrate", "carbs"),
    (r"^fib(er|re)|dietaryfib", "fiber"),
    (r"^sugar|totalsugar|sugarstotal", "sugar"),
    (r"saturat|^fasat|^sfa|satfat", "satFat"),
    (r"monounsat|^famono|^mufa", "mufa"),
    (r"polyunsat|^fapoly|^pufa", "pufa"),
    (r"18:3|linolenic|\bala\b|^ala", "ala"),
    (r"20:5|eicosapent|^epa", "epa"),
    (r"22:6|docosahex|^dha", "dha"),
    (r"18:2|linoleic|^la\b", "la"),
    (r"retinolactivity|vit(amin)?a.*rae|^vitarae", "vitA"),
    (r"^vit(amin)?a(iu|re|\b)", "vitA"),
    (r"^vit(amin)?c|ascorbic", "vitC"),
    (r"^vit(amin)?d", "vitD"),
    (r"^vit(amin)?e|tocopherol", "vitE"),
    (r"^vit(amin)?k|phylloquin", "vitK"),
    (r"thiamin|^vit(amin)?b1\b|^b1", "b1"),
    (r"riboflav|^vit(amin)?b2\b|^b2", "b2"),
    (r"niacin|^vit(amin)?b3\b", "b3"),
    (r"panto|^vit(amin)?b5\b", "b5"),
    (r"^vit(amin)?b6|pyridox", "b6"),
    (r"biotin|^vit(amin)?b7", "b7"),
    (r"folate|folic|^vit(amin)?b9", "folate"),
    (r"b12|cobalamin", "b12"),
    (r"choline", "choline"),
    (r"^calcium|^ca\b", "calcium"),
    (r"^iron|^fe\b", "iron"),
    (r"^magnes|^mg\b", "magnesium"),
    (r"^phosph|^p\b", "phosphorus"),
    (r"^potass|^k\b", "potassium"),
    (r"^zinc|^zn\b", "zinc"),
    (r"^copper|^cu\b", "copper"),
    (r"^mangan|^mn\b", "manganese"),
    (r"^seleni|^se\b", "selenium"),
    (r"^iodin|^i\b", "iodine"),
    (r"^sodium|^na\b", "sodium"),
]


def guess_unit(col: str) -> str | None:
    c = str(col).lower().replace("μ", "µ")
    for pat, unit in [(r"\brae\b|[_(]rae\b|\(rae", "µg RAE"), (r"dfe", "µg DFE"), (r"kcal", "kcal"), (r"\bkj\b|\(kj\)|_kj", "kJ"), (r"µg|\bug\b|mcg|_ug", "µg"), (r"\(mg\)|_mg|\bmg\b", "mg"),
                      (r"\(g\)|_g\b|\bg\b", "g"), (r"\biu\b|_iu", "IU")]:
        if re.search(pat, c):
            return unit
    return None


def suggest_mapping(columns: list[Any]) -> dict[str, list[dict[str, str]]]:
    out: dict[str, list[dict[str, str]]] = {}
    for col in columns:
        nc = norm_col(col)
        for pat, key in SUGGEST:
            if re.search(pat, nc):
                spec = {"column": str(col)}
                u = guess_unit(col)
                if u:
                    spec["unit"] = u
                out.setdefault(key, []).append(spec)
                break
    return out


def guess_header_row(rows: list[tuple]) -> int:
    best, best_n = 0, -1
    for i, r in enumerate(rows[:10]):
        cells = [c for c in r if c not in (None, "")]
        strs = [c for c in cells if isinstance(c, str)]
        if len(cells) >= 3 and len(strs) >= 0.6 * len(cells) and len(strs) > best_n:
            best, best_n = i, len(strs)
    return best


def inspect_workbook(path: str, cfg: dict | None = None, n_rows: int = 3, count_rows: bool = True,
                     out=print) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    report = []
    configured = {norm_col(s["sheet"]): s["key"] for s in (cfg or {}).get("sources", [])}
    for ws in wb.worksheets:
        head: list[tuple] = []
        nrows = 0
        ncols = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 12:
                head.append(row)
            if i < 12 or count_rows:
                nrows = i + 1
                ncols = max(ncols, len(row))
            if not count_rows and i >= 12:
                break
        hdr = guess_header_row(head)
        columns = [c for c in (head[hdr] if head else ()) if c not in (None, "")]
        out("=" * 100)
        tag = f"   [configured as source '{configured[norm_col(ws.title)]}']" if norm_col(ws.title) in configured else ""
        out(f"SHEET {ws.title!r}: {nrows if count_rows else '>=' + str(nrows)} rows x {ncols} cols{tag}")
        out(f"  header row guess (0-based): {hdr}")
        for i, r in enumerate(head[: hdr + 1 + n_rows]):
            vals = [("" if c is None else str(c))[:24] for c in r[:14]]
            out(f"  row {i}: {vals}{' ...' if len(r) > 14 else ''}")
        out(f"  columns ({len(columns)}): {columns}")
        sug = suggest_mapping(columns)
        if sug:
            out("  suggested nutrient mapping (verify!):")
            out("  " + json.dumps(sug, ensure_ascii=False)[:2500])
        report.append({"sheet": ws.title, "rows": nrows, "cols": ncols, "header_row": hdr, "columns": columns,
                       "suggested": sug})
    wb.close()
    if cfg:
        have = {norm_col(r["sheet"]) for r in report}
        for s in cfg.get("sources", []):
            if s.get("enabled", True) and norm_col(s["sheet"]) not in have:
                out(f"!! configured source {s['key']!r} expects sheet {s['sheet']!r} which is NOT in the workbook")
    return report
